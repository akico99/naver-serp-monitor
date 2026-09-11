# -*- coding: utf-8 -*-
"""data/키워드_통합정리.xlsx 의 '통합_키워드' 시트를 웹에서 바로 읽을 수 있는
data/master_keywords.js (window.MASTER_KEYWORDS = [...]) 로 변환한다.

원본 엑셀은 사람이 수기로 관리하므로, 이 스크립트는 실행할 때마다 최신 엑셀 내용으로
master_keywords.js를 다시 만든다 (엑셀이 바뀌면 재실행하면 됨).

사용례:
    python collector/xlsx_to_master.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

PROJECT_ROOT = Path(__file__).parent.parent
XLSX_PATH = PROJECT_ROOT / "data" / "키워드_통합정리.xlsx"
OUT_PATH = PROJECT_ROOT / "data" / "master_keywords.js"

NICK_COLS = ["팔레트", "로지스트", "토배기", "레온", "무주스키장", "타블로그"]


def _slug(text: str) -> str:
    return "".join(ch if ch.isalnum() else "_" for ch in text).strip("_")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["통합_키워드"]
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    headers = [c.value for c in header_row]
    idx = {h: i for i, h in enumerate(headers) if h}

    def get(row, name):
        i = idx.get(name)
        return row[i] if i is not None else None

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        keyword = get(row, "키워드")
        if not keyword or not str(keyword).strip():
            continue
        keyword = str(keyword).strip()

        assignments = {}
        for col in NICK_COLS:
            v = get(row, col)
            if v not in (None, ""):
                assignments[col] = str(v).strip()

        volume = get(row, "평균검색량")
        try:
            volume = float(volume) if volume not in (None, "") else None
        except (TypeError, ValueError):
            volume = None

        rows.append(
            {
                "출처": get(row, "출처"),
                "분류": get(row, "분류"),
                "대표키워드": get(row, "대표키워드"),
                "키워드": keyword,
                "행유형": get(row, "행유형"),
                "평균검색량": volume,
                "담당": assignments,
                "비고": get(row, "비고"),
                "slug": _slug(keyword),
            }
        )

    OUT_PATH.write_text(
        "// 자동 생성 파일: collector/xlsx_to_master.py 가 data/키워드_통합정리.xlsx 에서 변환합니다. 직접 수정하지 마세요.\n"
        f"window.MASTER_KEYWORDS = {json.dumps(rows, ensure_ascii=False, indent=2)};\n",
        encoding="utf-8",
    )
    print(f"{len(rows)}개 키워드 -> {OUT_PATH}")


if __name__ == "__main__":
    main()
